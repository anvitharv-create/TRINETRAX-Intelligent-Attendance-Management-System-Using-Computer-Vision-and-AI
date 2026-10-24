"""Excel export for the Principal/Admin reports page.

The export endpoint hands this module the *same* payloads the on-screen
endpoints return, so a workbook can never disagree with what the user saw.
Uses openpyxl - the same library the HOD report already uses.
"""
import io
from datetime import datetime

HEADER_FONT = {"bold": True, "color": "FFFFFF"}
HEADER_FILL = "0F172A"


def _write_sheet(wb, title, headers, rows, first=False):
    from openpyxl.styles import Font, PatternFill
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color=HEADER_FONT["color"])
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
    for row in rows:
        ws.append(row)
    for i, _ in enumerate(headers, start=1):
        width = max([len(str(headers[i - 1]))] + [len(str(r[i - 1])) for r in rows if r]) if rows else len(str(headers[i - 1]))
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(max(width + 2, 10), 40)
    return ws


def _meta_rows(meta):
    f = meta.get("filters", {})
    return [
        ["Generated on", meta.get("generated_at")],
        ["Report", meta.get("title", "")],
        ["Date range", "%s to %s" % (f.get("start_date") or "beginning", f.get("end_date") or "today")],
        ["Department", f.get("branch") or "all"],
        ["Semester", f.get("semester") or "all"],
        ["Section", f.get("section") or "all"],
        ["Subject", f.get("subject") or "all"],
        ["Attendance threshold %", meta.get("threshold_pct")],
    ]


def render_xlsx(report_type, payload, meta) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    meta_ws = wb.active
    meta_ws.title = "Report"
    meta_ws["A1"] = "TrinetraX - %s" % meta.get("title", "Report")
    meta_ws["A1"].font = Font(bold=True, size=14)
    for i, (label, value) in enumerate(_meta_rows(meta), start=3):
        meta_ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        meta_ws.cell(row=i, column=2, value=value)
    meta_ws.column_dimensions["A"].width = 22
    meta_ws.column_dimensions["B"].width = 46

    if report_type == "college":
        ov = payload
        summary = [
            ["Total students", ov["total_students"]],
            ["Total teachers", ov["total_teachers"]],
            ["Departments", ov["total_departments"]],
            ["Overall attendance %", ov["overall_attendance_pct"]],
            ["Records in window", ov["records_in_window"]],
            ["Present today", ov["present_today"]],
            ["Absent today", ov["absent_today"]],
            ["Students below threshold", ov["low_attendance_count"]],
        ]
        _write_sheet(wb, "Summary", ["Metric", "Value"], summary)
        _write_sheet(
            wb, "Departments",
            ["Department", "Students", "Teachers", "Subjects", "Attendance %", "Records", "Below threshold"],
            [[b["branch"], b["students"], b["teachers"], b["subjects"], b["attendance_pct"],
              b["records"], b["below_threshold_students"]] for b in ov.get("branch_stats", [])],
        )
        _write_sheet(
            wb, "Trend",
            ["Date", "Attendance %", "Records"],
            [[t["date"], t["pct"], t["records"]] for t in ov.get("monthly_trend", [])],
        )
    elif report_type == "department":
        dept = payload.get("department") or {}
        _write_sheet(
            wb, "Semesters", ["Semester", "Students", "Attendance %"],
            [[s["semester"], s["students"], s["attendance_pct"]] for s in dept.get("semesters", [])],
            first=False,
        )
        sem = payload.get("semester")
        if sem:
            _write_sheet(
                wb, "Sections", ["Section", "Students", "Attendance %"],
                [[s["section"], s["students"], s["attendance_pct"]] for s in sem.get("sections", [])],
            )
        cls = payload.get("class")
        if cls:
            _write_sheet(
                wb, "Students",
                ["USN", "Name", "Roll No", "Attendance %", "Today"],
                [[s["student_id"], s["name"], s["roll_number"], s["attendance_pct"], s["today_status"]]
                 for s in cls.get("students", [])],
            )
    elif report_type == "subject":
        _write_sheet(
            wb, "Students",
            ["USN", "Name", "Total classes", "Present", "Absent", "Other", "Attendance %"],
            [[r["usn"], r["name"], r["total"], r["present"], r["absent"], r["other"], r["pct"]]
             for r in payload],
        )
    elif report_type == "low":
        has_sem = bool(payload) and "semester" in payload[0]
        head = ["USN", "Name", "Roll No", "Branch"] + (["Semester", "Section"] if has_sem else []) + ["Attendance %"]
        rows = []
        for s in payload:
            row = [s.get("student_id"), s.get("name"), s.get("roll_number"), s.get("branch")]
            if has_sem:
                row += [s.get("semester"), s.get("section")]
            row.append(s.get("percentage"))
            rows.append(row)
        _write_sheet(wb, "Students", head, rows)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
